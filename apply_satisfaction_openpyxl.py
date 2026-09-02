# -*- coding: utf-8 -*-
"""
회차별 만족도를 원본 엑셀에 셀 단위로 써넣는다 (서식 유지).

pandas 로 다시 쓰면 원본 서식과 압축이 날아가 15MB -> 53MB 로 불어난다.
openpyxl 로 값 셀만 갈아끼우면 원본 크기가 유지된다.

건드리는 곳
-----------
  Sheet1                      '만족도' 컬럼  <- 회차별 5점 실측
                              + 평가인원/평가참여율/추천인원/추천응답인원/설문대상인원
                              + 만족도_미집계사유
  k_digital_training_1_260805 L열 '만족도 점수' <- 같은 값

두 시트의 만족도는 같은 값이어야 한다. 원래 L열에는 오픈 API 의 stdgScor(100점,
과정 단위 통합값)가 들어 있었는데, 그 값은 같은 과정의 모든 회차가 동일해서
회차별 편차를 통째로 가린다. 회차별 실측으로 대체한다.

조인 키
-------
  Sheet1                      : '훈련과정 ID' + '회차'
  k_digital_training_1_260805 : '훈련과정ID'(Z) + '훈련과정 순차'(Y)
두 시트 모두 7,230행으로 1:1 대응한다.
"""

import argparse
import csv
import os
import shutil
import sys
from datetime import datetime

from openpyxl import load_workbook

TODAY = datetime(2026, 9, 2)

MAIN_SHEET = "Sheet1"
SRC_SHEET = "k_digital_training_1_260805"

# Sheet1 에 새로 붙일 컬럼 (CSV 컬럼명 -> 엑셀 헤더)
EXTRA = {
    "평가인원": "평가인원",
    "평가참여율": "평가참여율",
    "추천인원": "추천인원",
    "추천응답인원": "추천응답인원",
    "수강인원": "설문대상인원",
}
REASON_COL = "만족도_미집계사유"


def header_map(ws):
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def to_num(v):
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load_sat(csv_path):
    """(훈련과정 ID, 회차) -> row"""
    out = {}
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            cid = str(r["훈련과정 ID"]).strip()
            tme = str(int(float(r["회차"])))
            out[(cid, tme)] = r
    return out


def classify(enrolled, end_dt, has_score):
    """만족도가 빈 이유. '값 없음'과 '회차 자체가 없음'을 구분하기 위한 것."""
    if has_score:
        return None
    if not enrolled or enrolled <= 0:
        return "미개설(수강신청 0)"
    if end_dt is None:
        return "설문 미실시"
    if end_dt > TODAY:
        return "진행중(종료 전)"
    if (TODAY - end_dt).days < 21:
        return "종료 직후(집계 전)"
    return "설문 미실시"


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--xlsx", required=True, help="원본 (수정하지 않음)")
    p.add_argument("--csv", required=True)
    p.add_argument("--output", default="")
    a = p.parse_args()

    out = a.output or "%s_회차별만족도.xlsx" % os.path.splitext(a.xlsx)[0]
    if os.path.abspath(out) == os.path.abspath(a.xlsx):
        print("원본과 출력 경로가 같습니다.", file=sys.stderr)
        sys.exit(1)

    sat = load_sat(a.csv)
    print("CSV: %d건" % len(sat))

    shutil.copy2(a.xlsx, out)          # 서식째 복사한 뒤 값만 고친다
    wb = load_workbook(out)

    # ── Sheet1 ──
    ws = wb[MAIN_SHEET]
    h = header_map(ws)
    for need in ("훈련과정 ID", "회차", "만족도", "수강신청 인원", "과정종료일"):
        if need not in h:
            print("Sheet1 에 '%s' 컬럼 없음" % need, file=sys.stderr)
            sys.exit(1)

    next_col = ws.max_column + 1
    newcols = {}
    for src, name in list(EXTRA.items()) + [(REASON_COL, REASON_COL)]:
        if name in h:
            newcols[src] = h[name]
        else:
            ws.cell(row=1, column=next_col, value=name)
            newcols[src] = next_col
            next_col += 1

    hit = miss = 0
    reasons = {}
    for row in range(2, ws.max_row + 1):
        cid = ws.cell(row=row, column=h["훈련과정 ID"]).value
        tme = ws.cell(row=row, column=h["회차"]).value
        if cid is None or tme is None:
            continue
        key = (str(cid).strip(), str(int(float(tme))))
        rec = sat.get(key)

        score = to_num(rec["만족도(5점)"]) if rec else None
        ws.cell(row=row, column=h["만족도"], value=score)

        for src, col in newcols.items():
            if src == REASON_COL:
                continue
            ws.cell(row=row, column=col, value=to_num(rec[src]) if rec else None)

        enrolled = to_num(ws.cell(row=row, column=h["수강신청 인원"]).value)
        end_val = ws.cell(row=row, column=h["과정종료일"]).value
        end_dt = end_val if isinstance(end_val, datetime) else None
        reason = classify(enrolled, end_dt, score is not None)
        ws.cell(row=row, column=newcols[REASON_COL], value=reason)

        if score is not None:
            hit += 1
        else:
            miss += 1
            reasons[reason] = reasons.get(reason, 0) + 1

    print("Sheet1: 실측 %d행 / 빈칸 %d행" % (hit, miss))
    for k, v in sorted(reasons.items(), key=lambda x: -x[1]):
        print("    %-20s %4d" % (k, v))

    # ── k_digital_training_1_260805 L열 ──
    ws2 = wb[SRC_SHEET]
    h2 = header_map(ws2)
    id_col = h2.get("훈련과정ID")
    seq_col = h2.get("훈련과정 순차")
    sat_col = h2.get("만족도 점수")
    if not all((id_col, seq_col, sat_col)):
        print("%s 에서 컬럼을 못 찾음: %s" % (SRC_SHEET, list(h2)[:10]), file=sys.stderr)
        sys.exit(1)

    hit2 = 0
    for row in range(2, ws2.max_row + 1):
        cid = ws2.cell(row=row, column=id_col).value
        seq = ws2.cell(row=row, column=seq_col).value
        if cid is None or seq is None:
            continue
        rec = sat.get((str(cid).strip(), str(int(float(seq)))))
        score = to_num(rec["만족도(5점)"]) if rec else None
        ws2.cell(row=row, column=sat_col, value=score)
        if score is not None:
            hit2 += 1

    print("%s L열('만족도 점수'): 실측 %d행" % (SRC_SHEET, hit2))

    wb.save(out)
    print("저장: %s (%.1f MB)" % (out, os.path.getsize(out) / 1024 / 1024))


if __name__ == "__main__":
    main()
